import os
import io
import csv
import re
from typing import List, Dict, Any, Tuple
import pypdf
from docx import Document

class DocumentParser:
    """
    Robust multi-format document parser supporting PDF, DOCX, TXT, CSV.
    Extracts raw text, lines, or tabular rows with fallback encodings.
    """

    @staticmethod
    def extract_text_from_bytes(file_bytes: bytes, filename: str) -> Tuple[str, str, List[Dict[str, Any]]]:
        """
        Returns (raw_text, file_type, tabular_rows_if_any)
        """
        ext = os.path.splitext(filename)[1].lower()
        file_type = ext.replace(".", "") or "txt"
        
        tabular_rows: List[Dict[str, Any]] = []
        raw_text = ""

        if ext == ".pdf":
            try:
                reader = pypdf.PdfReader(io.BytesIO(file_bytes))
                text_pages = []
                for idx, page in enumerate(reader.pages):
                    extracted = page.extract_text()
                    if extracted:
                        text_pages.append(extracted)
                raw_text = "\n".join(text_pages)
            except Exception as e:
                raw_text = f"Error reading PDF: {str(e)}"
                
        elif ext in [".docx", ".doc"]:
            try:
                doc = Document(io.BytesIO(file_bytes))
                paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
                # Also extract table cells if any
                table_texts = []
                for table in doc.tables:
                    for row in table.rows:
                        row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                        if row_cells:
                            table_texts.append(" | ".join(row_cells))
                raw_text = "\n".join(paragraphs + table_texts)
            except Exception as e:
                raw_text = f"Error reading DOCX: {str(e)}"
                
        elif ext in [".csv", ".tsv"]:
            try:
                text_decoded = DocumentParser._decode_bytes(file_bytes)
                delimiter = "\t" if ext == ".tsv" or "\t" in text_decoded.split("\n")[0] else ","
                reader = csv.DictReader(io.StringIO(text_decoded), delimiter=delimiter)
                for row in reader:
                    cleaned_row = {k.strip().lower(): v.strip() for k, v in row.items() if k}
                    tabular_rows.append(cleaned_row)
                raw_text = text_decoded
            except Exception as e:
                raw_text = DocumentParser._decode_bytes(file_bytes)
                
        else: # .txt or generic
            raw_text = DocumentParser._decode_bytes(file_bytes)

        return raw_text, file_type, tabular_rows

    @staticmethod
    def _decode_bytes(content: bytes) -> str:
        for enc in ["utf-8", "utf-8-sig", "latin-1", "cp1252", "iso-8859-1"]:
            try:
                return content.decode(enc)
            except UnicodeDecodeError:
                continue
        return content.decode("utf-8", errors="ignore")
